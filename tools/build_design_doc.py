"""Build the design document as a Word file from docs/design-doc.md.

    py tools/build_design_doc.py            -> docs/The Fourth Estate - Game Design Document.docx
    py tools/build_design_doc.py out.docx

The prose lives in docs/design-doc.md (a small Markdown subset: # headings,
paragraphs with **bold** / *italic*, "- " bullets, "- [ ] " checkboxes,
"1. " numbered steps, pipe tables, and ![caption](path) images). Numbers and
both appendices come from backend/game_data.php through the simulator's own
parser, so the document can never disagree with the game:

    {{n_cards}} {{n_event}} {{n_profit}} {{n_opening}} {{n_races}}
    {{n_races_word}} {{n_races_word_cap}} {{n_nation_wins}} {{n_states_wins}}
    {{example_rows}} {{appendix_races}} {{appendix_cards}}

Tables follow the house style: a thin black grid, a light grey header row,
padded cells. Opens directly in Google Docs (File > Open > Upload).
"""
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import simulate as sim  # noqa: E402  (parses game_data.php and checks it)

SOURCE = os.path.join(ROOT, "docs", "design-doc.md")
DEFAULT_OUT = os.path.join(ROOT, "docs", "The Fourth Estate - Game Design Document.docx")
EXAMPLES = ["bill_of_rights", "federalist_papers", "kansas_nebraska", "postal_act_1792"]
WORDS = {14: "fourteen", 15: "fifteen", 16: "sixteen", 17: "seventeen", 18: "eighteen"}

# ---------------------------------------------------------------- content from the game


def push(v):
    if v == 0:
        return "—"
    return ("Nation +%d" if v > 0 else "States +%d") % abs(v)


def story_row(c, enters=None):
    trade = c["kind"] == "profit"
    cells = [c["name"] + (" (trade)" if trade else ""), str(c["year"])]
    if enters is not None:
        cells.append(enters)
    cells += [str(c["profit"]), push(c["positive"]), push(c["negative"]), str(c["stability"] or "—")]
    return cells


def table_md(header, rows):
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join("---" for _ in header) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in rows]
    return "\n".join(lines)


def game_facts():
    cards, races = sim.CARDS, sim.ELECTIONS
    years = [e["year"] for e in races]

    def enters(y):
        first = next(ey for ey in years if y <= ey)
        return "opening deck" if first == years[0] else str(first)

    nation = sum(1 for e in races if e["historical_winner"] == "nation")
    n = len(races)
    race_rows = []
    for e in races:
        w = e[e["historical_winner"]]
        surname = "J.Q. Adams" if w["name"] == "John Quincy Adams" else w["name"].split()[-1]
        race_rows.append([str(e["space"]), str(e["year"]),
                          "%s (%s)" % (e["nation"]["name"], e["nation"]["party"]),
                          "%s (%s)" % (e["states"]["name"], e["states"]["party"]),
                          "%s — %s" % (surname, e["historical_winner"].capitalize())])
    ordered = sorted(cards.values(), key=lambda c: c["year"])
    return {
        "n_cards": str(len(cards)),
        "n_event": str(sum(1 for c in cards.values() if c["kind"] == "event")),
        "n_profit": str(sum(1 for c in cards.values() if c["kind"] == "profit")),
        "n_opening": str(len(sim.OPENING)),
        "n_races": str(n),
        "n_races_word": WORDS.get(n, str(n)),
        "n_races_word_cap": WORDS.get(n, str(n)).capitalize(),
        "n_nation_wins": str(nation),
        "n_states_wins": str(n - nation),
        "example_rows": "\n".join("| " + " | ".join(story_row(cards[k])) + " |" for k in EXAMPLES),
        "appendix_races": table_md(["#", "Year", "Nation (federal power)", "States (states' rights)", "History elected"],
                                   race_rows),
        "appendix_cards": table_md(["Story", "Year", "Enters", "Profit (bury)", "Positive story", "Negative story",
                                    "Stability cost"],
                                   [story_row(c, enters(c["year"])) for c in ordered]),
    }


# ---------------------------------------------------------------- Word formatting


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def style_table(table):
    """House style: thin black grid, padded cells."""
    tbl_pr = table._tbl.tblPr
    for old in tbl_pr.findall(qn("w:tblW")):
        tbl_pr.remove(old)
    width = OxmlElement("w:tblW")          # full text width
    width.set(qn("w:w"), "5000")
    width.set(qn("w:type"), "pct")
    tbl_pr.insert(0, width)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for edge, w in (("top", 60), ("left", 100), ("bottom", 60), ("right", 100)):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:w"), str(w))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)
    # Word insists on the schema's child order inside tblPr.
    order = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
             "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]
    kids = list(tbl_pr)
    for k in kids:
        tbl_pr.remove(k)
    kids.sort(key=lambda k: order.index(k.tag.split("}")[1]) if k.tag.split("}")[1] in order else len(order))
    for k in kids:
        tbl_pr.append(k)


def add_runs(par, text):
    """**bold** and *italic* inline."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("*"):
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def add_table(doc, lines):
    rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines]
    header, body = rows[0], [r for r in rows[2:]]
    table = doc.add_table(rows=1 + len(body), cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    style_table(table)
    # Column widths follow the text: long prose columns get the room.
    ncol = len(header)
    page = 6.5                                  # inches of text width
    floor, share = [], []
    for ci in range(ncol):
        texts = [r[ci] for r in [header] + body if ci < len(r)]
        longest_word = max((len(w) for t in texts for w in t.split()), default=1)
        floor.append(0.095 * longest_word + 0.25)   # never split a word (bold 11pt Arial + padding)
        share.append(min(60, sum(len(t) for t in texts) / len(texts)))
    inches = [max(f, page * s / sum(share)) for f, s in zip(floor, share)]
    over = sum(inches) - page
    if over > 0:                                # take the excess from columns above their floor
        slack = [i - f for i, f in zip(inches, floor)]
        inches = [i - over * sl / sum(slack) for i, sl in zip(inches, slack)]
    widths = [Inches(i) for i in inches]
    for ci, col in enumerate(table.columns):
        col.width = widths[ci]
    for ri, r in enumerate([header] + body):
        tr_pr = table.rows[ri]._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")    # never break a row across pages
        cant.set(qn("w:val"), "true")
        tr_pr.append(cant)
        for ci in range(ncol):
            table.cell(ri, ci).width = widths[ci]
        if ri == 0:  # repeat the header row on every page
            h = OxmlElement("w:tblHeader")
            h.set(qn("w:val"), "true")
            tr_pr.append(h)
        for ci in range(len(header)):
            cell = table.cell(ri, ci)
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            add_runs(par, r[ci] if ci < len(r) else "")
            if ri == 0:
                set_cell_shading(cell, "F2F2F2")
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.keep_with_next = True   # no header orphaned at a page foot
                for run in par.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in (("Heading 1", 20), ("Heading 2", 15)):
        st = doc.styles[name]
        st.font.name = "Arial"
        fonts = st.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fonts.attrib.pop(qn(attr), None)   # theme fonts would override Arial
        fonts.set(qn("w:eastAsia"), "Arial")
        st.font.size = Pt(size)
        st.font.bold = name == "Heading 2"
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(18 if name == "Heading 2" else 0)
        st.paragraph_format.space_after = Pt(6)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(1)


def build(out_path):
    text = open(SOURCE, encoding="utf-8").read()
    for key, val in game_facts().items():
        text = text.replace("{{%s}}" % key, val)
    missing = re.findall(r"\{\{\w+\}\}", text)
    assert not missing, "unknown placeholders: %s" % missing

    doc = Document()
    setup_styles(doc)
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            i += 1
            continue
        if s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        if s.startswith("# "):
            doc.add_heading(s[2:], level=1)
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=2)
        elif s.startswith("!["):
            m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", s)
            doc.add_picture(os.path.join(ROOT, "docs", m.group(2)), width=Inches(6))
        elif s.startswith("- [ ] "):
            add_runs(doc.add_paragraph(style="List Bullet"), "☐ " + s[6:])
        elif s.startswith("- "):
            add_runs(doc.add_paragraph(style="List Bullet"), s[2:])
        elif re.match(r"\d+\. ", s):
            add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", s))
        else:
            add_runs(doc.add_paragraph(), s)
        i += 1
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    print(build(sys.argv[1] if len(sys.argv) > 1 else DEFAULT_OUT))
